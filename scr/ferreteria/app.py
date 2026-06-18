from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
import mysql.connector
from config import DB_CONFIG

from datetime import datetime
from io import BytesIO, StringIO
from decimal import Decimal
from collections import defaultdict
from email.message import EmailMessage
from ipc_routes import ipc_bp

import csv
import os
import re
import smtplib

import openpyxl
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.cell.cell import MergedCell


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, '.env'))

app = Flask(__name__)
app.secret_key = 'ferreteria_secret_key'

app.register_blueprint(ipc_bp)

UPLOAD_FOLDER = os.path.join(BASE_DIR, 'uploads')
LOG_FOLDER = os.path.join(BASE_DIR, 'logs')

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(LOG_FOLDER, exist_ok=True)


def get_db_connection():
    return mysql.connector.connect(**DB_CONFIG)


# =========================================================
# EXPORTACIÓN ESTADÍSTICA PROFESIONAL DE MOVIMIENTOS
# =========================================================

COLUMNAS_EXPORTACION = [
    'id_movimiento',
    'fecha',
    'tipo_movimiento',
    'id_producto',
    'producto',
    'categoria',
    'stock_anterior',
    'cantidad',
    'stock_posterior',
    'precio_unitario',
    'importe_estimado',
    'usuario_responsable',
    'motivo'
]


ENCABEZADOS_EXCEL = [
    'ID Movimiento',
    'Fecha',
    'Tipo de Movimiento',
    'ID Producto',
    'Producto',
    'Categoría',
    'Stock Anterior',
    'Cantidad',
    'Stock Posterior',
    'Precio Unitario',
    'Importe Estimado',
    'Usuario Responsable',
    'Motivo'
]


def usuario_puede_exportar():
    return session.get('usuario_rol') in ['encargado', 'admin']


def convertir_valor_csv(valor):
    if valor is None:
        return ''

    if isinstance(valor, Decimal):
        return float(valor)

    if isinstance(valor, datetime):
        return valor.strftime('%Y-%m-%d %H:%M:%S')

    return valor


def convertir_valor_excel(valor):
    if valor is None:
        return ''

    if isinstance(valor, Decimal):
        return float(valor)

    return valor


def validar_rango_fechas(fecha_desde, fecha_hasta):
    if not fecha_desde or not fecha_hasta:
        return False, 'Debe seleccionar fecha desde y fecha hasta.'

    try:
        desde = datetime.strptime(fecha_desde, '%Y-%m-%d').date()
        hasta = datetime.strptime(fecha_hasta, '%Y-%m-%d').date()
    except ValueError:
        return False, 'El formato de fecha no es válido.'

    if desde > hasta:
        return False, 'La fecha inicial no puede ser mayor que la fecha final.'

    return True, ''


def consultar_movimientos_exportacion(fecha_desde, fecha_hasta, tipo_movimiento='todos', id_categoria=''):
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    query = """
        SELECT
            m.id AS id_movimiento,
            m.fecha AS fecha,
            m.tipo AS tipo_movimiento,
            p.id AS id_producto,
            p.nombre AS producto,
            COALESCE(c.nombre, 'Sin categoría') AS categoria,
            m.stock_anterior AS stock_anterior,
            m.cantidad AS cantidad,
            m.stock_posterior AS stock_posterior,
            COALESCE(m.precio_unitario, p.precio) AS precio_unitario,
            CASE
                WHEN m.tipo IN ('entrada', 'salida')
                    THEN m.cantidad * COALESCE(m.precio_unitario, p.precio)
                ELSE 0
            END AS importe_estimado,
            u.nombre AS usuario_responsable,
            COALESCE(m.motivo, '') AS motivo
        FROM movimientos_stock m
        JOIN productos p ON m.id_producto = p.id
        LEFT JOIN categorias c ON p.id_categoria = c.id
        JOIN usuarios u ON m.id_usuario = u.id
        WHERE DATE(m.fecha) BETWEEN %s AND %s
    """

    params = [fecha_desde, fecha_hasta]

    if tipo_movimiento and tipo_movimiento != 'todos':
        query += " AND m.tipo = %s"
        params.append(tipo_movimiento)

    if id_categoria:
        query += " AND p.id_categoria = %s"
        params.append(id_categoria)

    query += " ORDER BY m.fecha ASC, m.id ASC"

    cursor.execute(query, params)
    movimientos = cursor.fetchall()

    cursor.close()
    conn.close()

    return movimientos


def registrar_auditoria_exportacion(fecha_desde, fecha_hasta, tipo_movimiento, id_categoria, formato, cantidad_registros):
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO exportaciones_estadisticas
            (id_usuario, fecha_desde, fecha_hasta, tipo_movimiento, id_categoria, formato, cantidad_registros)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """, (
            session['usuario_id'],
            fecha_desde,
            fecha_hasta,
            tipo_movimiento or 'todos',
            id_categoria if id_categoria else None,
            formato,
            cantidad_registros
        ))

        conn.commit()
        cursor.close()
        conn.close()

    except Exception as error:
        print('No se pudo registrar auditoría de exportación:', error)


def generar_csv_movimientos(movimientos, fecha_desde, fecha_hasta):
    salida = StringIO()

    writer = csv.DictWriter(
        salida,
        fieldnames=COLUMNAS_EXPORTACION,
        delimiter=';',
        lineterminator='\n'
    )

    writer.writeheader()

    for movimiento in movimientos:
        fila = {}

        for columna in COLUMNAS_EXPORTACION:
            fila[columna] = convertir_valor_csv(movimiento.get(columna))

        writer.writerow(fila)

    contenido = salida.getvalue().encode('utf-8-sig')
    archivo = BytesIO(contenido)
    archivo.seek(0)

    nombre_archivo = f'movimientos_stock_estadistica_{fecha_desde}_a_{fecha_hasta}.csv'

    return send_file(
        archivo,
        mimetype='text/csv; charset=utf-8',
        as_attachment=True,
        download_name=nombre_archivo
    )


def normalizar_nombre_tabla(nombre):
    nombre_limpio = re.sub(r'[^A-Za-z0-9_]', '_', nombre)
    return nombre_limpio[:50]


def aplicar_bordes_y_alineacion(ws, min_row, max_row, min_col, max_col):
    borde_suave = Border(
        left=Side(style='thin', color='D8DEE9'),
        right=Side(style='thin', color='D8DEE9'),
        top=Side(style='thin', color='D8DEE9'),
        bottom=Side(style='thin', color='D8DEE9')
    )

    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue

            cell.border = borde_suave
            cell.alignment = Alignment(vertical='center', wrap_text=True)


def configurar_impresion(ws):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.50
    ws.page_margins.bottom = 0.50
    ws.sheet_view.showGridLines = False


def ajustar_columnas_detalle(ws):
    anchos = {
        'A': 14,
        'B': 20,
        'C': 20,
        'D': 14,
        'E': 32,
        'F': 24,
        'G': 16,
        'H': 12,
        'I': 16,
        'J': 18,
        'K': 18,
        'L': 26,
        'M': 48
    }

    for columna, ancho in anchos.items():
        ws.column_dimensions[columna].width = ancho


def escribir_portada_hoja(ws, titulo, subtitulo, fecha_desde, fecha_hasta):
    ws.merge_cells('A1:M1')
    ws['A1'] = titulo
    ws['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='17324D', end_color='17324D', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    ws.merge_cells('A2:M2')
    ws['A2'] = subtitulo
    ws['A2'].font = Font(name='Calibri', size=11, italic=True, color='334155')
    ws['A2'].fill = PatternFill(start_color='EAF2FF', end_color='EAF2FF', fill_type='solid')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 24

    ws.merge_cells('A3:M3')
    ws['A3'] = (
        f'Período analizado: {fecha_desde} a {fecha_hasta}   |   '
        f'Generado por: {session.get("usuario_nombre", "Usuario no identificado")}'
    )
    ws['A3'].font = Font(name='Calibri', size=10, color='475569')
    ws['A3'].fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    ws['A3'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[3].height = 22


def agregar_tabla_excel(ws, nombre_tabla, fila_inicio, columna_inicio, fila_fin, columna_fin):
    if fila_fin <= fila_inicio:
        return

    referencia = f'{columna_inicio}{fila_inicio}:{columna_fin}{fila_fin}'
    tabla = Table(displayName=normalizar_nombre_tabla(nombre_tabla), ref=referencia)

    estilo = TableStyleInfo(
        name='TableStyleMedium9',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    tabla.tableStyleInfo = estilo
    ws.add_table(tabla)


def agregar_hoja_detalle(wb, nombre_hoja, movimientos, fecha_desde, fecha_hasta, color_pestana, descripcion):
    ws = wb.create_sheet(nombre_hoja)
    ws.sheet_properties.tabColor = color_pestana

    configurar_impresion(ws)

    escribir_portada_hoja(
        ws=ws,
        titulo=f'{nombre_hoja.upper()} — REPORTE DE MOVIMIENTOS',
        subtitulo=descripcion,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta
    )

    fila_encabezado = 5

    for col, encabezado in enumerate(ENCABEZADOS_EXCEL, start=1):
        celda = ws.cell(row=fila_encabezado, column=col)
        celda.value = encabezado
        celda.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
        celda.fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.row_dimensions[fila_encabezado].height = 28

    fila_actual = fila_encabezado + 1

    for movimiento in movimientos:
        valores = [
            convertir_valor_excel(movimiento.get('id_movimiento')),
            convertir_valor_excel(movimiento.get('fecha')),
            convertir_valor_excel(movimiento.get('tipo_movimiento')),
            convertir_valor_excel(movimiento.get('id_producto')),
            convertir_valor_excel(movimiento.get('producto')),
            convertir_valor_excel(movimiento.get('categoria')),
            convertir_valor_excel(movimiento.get('stock_anterior')),
            convertir_valor_excel(movimiento.get('cantidad')),
            convertir_valor_excel(movimiento.get('stock_posterior')),
            convertir_valor_excel(movimiento.get('precio_unitario')),
            convertir_valor_excel(movimiento.get('importe_estimado')),
            convertir_valor_excel(movimiento.get('usuario_responsable')),
            convertir_valor_excel(movimiento.get('motivo'))
        ]

        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila_actual, column=col)
            celda.value = valor
            celda.font = Font(name='Calibri', size=10, color='1E293B')
            celda.alignment = Alignment(vertical='center', wrap_text=True)

        tipo = movimiento.get('tipo_movimiento')
        celda_tipo = ws.cell(row=fila_actual, column=3)

        if tipo == 'entrada':
            celda_tipo.fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
            celda_tipo.font = Font(name='Calibri', size=10, bold=True, color='166534')

        elif tipo == 'salida':
            celda_tipo.fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
            celda_tipo.font = Font(name='Calibri', size=10, bold=True, color='991B1B')

        elif tipo == 'ajuste':
            celda_tipo.fill = PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid')
            celda_tipo.font = Font(name='Calibri', size=10, bold=True, color='1D4ED8')

        ws.cell(row=fila_actual, column=2).number_format = 'dd/mm/yyyy hh:mm'
        ws.cell(row=fila_actual, column=7).number_format = '0'
        ws.cell(row=fila_actual, column=8).number_format = '0'
        ws.cell(row=fila_actual, column=9).number_format = '0'
        ws.cell(row=fila_actual, column=10).number_format = '$ #,##0.00'
        ws.cell(row=fila_actual, column=11).number_format = '$ #,##0.00'

        fila_actual += 1

    if movimientos:
        fila_fin = fila_actual - 1

        aplicar_bordes_y_alineacion(
            ws=ws,
            min_row=fila_encabezado,
            max_row=fila_fin,
            min_col=1,
            max_col=13
        )

        agregar_tabla_excel(
            ws=ws,
            nombre_tabla=f'Tabla_{nombre_hoja}',
            fila_inicio=fila_encabezado,
            columna_inicio='A',
            fila_fin=fila_fin,
            columna_fin='M'
        )

    else:
        ws.merge_cells('A7:M7')
        ws['A7'] = 'No existen movimientos para esta sección.'
        ws['A7'].font = Font(name='Calibri', size=11, italic=True, color='64748B')
        ws['A7'].alignment = Alignment(horizontal='center', vertical='center')
        ws['A7'].fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    ws.freeze_panes = 'A6'
    ws.auto_filter.ref = f'A{fila_encabezado}:M{max(fila_encabezado, ws.max_row)}'

    ajustar_columnas_detalle(ws)

    return ws


def calcular_totales(movimientos):
    ventas = [m for m in movimientos if m.get('tipo_movimiento') == 'salida']
    compras = [m for m in movimientos if m.get('tipo_movimiento') == 'entrada']
    ajustes = [m for m in movimientos if m.get('tipo_movimiento') == 'ajuste']

    unidades_vendidas = sum(int(m.get('cantidad') or 0) for m in ventas)
    unidades_compradas = sum(int(m.get('cantidad') or 0) for m in compras)
    unidades_ajustadas = sum(int(m.get('cantidad') or 0) for m in ajustes)

    importe_ventas = sum(float(m.get('importe_estimado') or 0) for m in ventas)
    importe_compras = sum(float(m.get('importe_estimado') or 0) for m in compras)

    promedio_salida = unidades_vendidas / len(ventas) if ventas else 0
    promedio_entrada = unidades_compradas / len(compras) if compras else 0

    return {
        'ventas': ventas,
        'compras': compras,
        'ajustes': ajustes,
        'total_movimientos': len(movimientos),
        'total_ventas': len(ventas),
        'total_compras': len(compras),
        'total_ajustes': len(ajustes),
        'unidades_vendidas': unidades_vendidas,
        'unidades_compradas': unidades_compradas,
        'unidades_ajustadas': unidades_ajustadas,
        'importe_ventas': importe_ventas,
        'importe_compras': importe_compras,
        'promedio_salida': promedio_salida,
        'promedio_entrada': promedio_entrada
    }


def calcular_producto_mayor_salida(movimientos):
    acumulado = defaultdict(int)

    for movimiento in movimientos:
        if movimiento.get('tipo_movimiento') == 'salida':
            acumulado[movimiento.get('producto')] += int(movimiento.get('cantidad') or 0)

    if not acumulado:
        return 'Sin salidas registradas', 0

    producto = max(acumulado, key=acumulado.get)
    return producto, acumulado[producto]


def calcular_categoria_mayor_movimiento(movimientos):
    acumulado = defaultdict(int)

    for movimiento in movimientos:
        categoria = movimiento.get('categoria') or 'Sin categoría'
        acumulado[categoria] += int(movimiento.get('cantidad') or 0)

    if not acumulado:
        return 'Sin datos', 0

    categoria = max(acumulado, key=acumulado.get)
    return categoria, acumulado[categoria]


def top_productos_salida(movimientos, limite=5):
    acumulado = defaultdict(int)

    for movimiento in movimientos:
        if movimiento.get('tipo_movimiento') == 'salida':
            acumulado[movimiento.get('producto')] += int(movimiento.get('cantidad') or 0)

    ordenado = sorted(acumulado.items(), key=lambda item: item[1], reverse=True)
    return ordenado[:limite]


def escribir_kpi(ws, rango, titulo, valor, color_fondo, color_texto='FFFFFF'):
    ws.merge_cells(rango)
    celda = ws[rango.split(':')[0]]
    celda.value = f'{titulo}\n{valor}'
    celda.font = Font(name='Calibri', size=14, bold=True, color=color_texto)
    celda.fill = PatternFill(start_color=color_fondo, end_color=color_fondo, fill_type='solid')
    celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    fila_inicio = int(re.findall(r'\d+', rango.split(':')[0])[0])
    ws.row_dimensions[fila_inicio].height = 46


def crear_hoja_resumen_profesional(wb, movimientos, fecha_desde, fecha_hasta):
    totales = calcular_totales(movimientos)

    producto_mayor, cantidad_producto_mayor = calcular_producto_mayor_salida(movimientos)
    categoria_mayor, cantidad_categoria_mayor = calcular_categoria_mayor_movimiento(movimientos)

    ws = wb.create_sheet('Resumen Ejecutivo', 0)
    ws.sheet_properties.tabColor = '17324D'
    ws.sheet_view.showGridLines = False

    ws.merge_cells('A1:L1')
    ws['A1'] = 'STOCKEADO — REPORTE PROFESIONAL DE MOVIMIENTOS DE STOCK'
    ws['A1'].font = Font(name='Calibri', size=20, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='17324D', end_color='17324D', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36

    ws.merge_cells('A2:L2')
    ws['A2'] = (
        f'Período analizado: {fecha_desde} a {fecha_hasta}   |   '
        f'Generado por: {session.get("usuario_nombre", "Usuario no identificado")}   |   '
        f'Fecha de generación: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}'
    )
    ws['A2'].font = Font(name='Calibri', size=11, italic=True, color='334155')
    ws['A2'].fill = PatternFill(start_color='EAF2FF', end_color='EAF2FF', fill_type='solid')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 24

    escribir_kpi(ws, 'A4:C5', 'TOTAL MOVIMIENTOS', totales['total_movimientos'], '2563EB')
    escribir_kpi(ws, 'D4:F5', 'SALIDAS / VENTAS', totales['total_ventas'], 'DC2626')
    escribir_kpi(ws, 'G4:I5', 'ENTRADAS / COMPRAS', totales['total_compras'], '16A34A')
    escribir_kpi(ws, 'J4:L5', 'AJUSTES', totales['total_ajustes'], '7C3AED')

    escribir_kpi(ws, 'A7:C8', 'UNIDADES VENDIDAS', totales['unidades_vendidas'], '991B1B')
    escribir_kpi(ws, 'D7:F8', 'UNIDADES COMPRADAS', totales['unidades_compradas'], '166534')
    escribir_kpi(ws, 'G7:I8', 'IMPORTE VENTAS', f"$ {totales['importe_ventas']:,.2f}", '0F766E')
    escribir_kpi(ws, 'J7:L8', 'IMPORTE COMPRAS', f"$ {totales['importe_compras']:,.2f}", '92400E')

    ws.merge_cells('A10:F10')
    ws['A10'] = 'INDICADORES PRINCIPALES'
    ws['A10'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    ws['A10'].fill = PatternFill(start_color='17324D', end_color='17324D', fill_type='solid')
    ws['A10'].alignment = Alignment(horizontal='center')

    indicadores = [
        ['Indicador', 'Valor'],
        ['Promedio de unidades por salida', round(totales['promedio_salida'], 2)],
        ['Promedio de unidades por entrada', round(totales['promedio_entrada'], 2)],
        ['Producto con mayor salida', producto_mayor],
        ['Cantidad acumulada del producto con mayor salida', cantidad_producto_mayor],
        ['Categoría con mayor movimiento', categoria_mayor],
        ['Cantidad acumulada en esa categoría', cantidad_categoria_mayor]
    ]

    fila_inicio = 11

    for i, fila in enumerate(indicadores, start=fila_inicio):
        ws.cell(row=i, column=1).value = fila[0]
        ws.cell(row=i, column=2).value = fila[1]

    aplicar_bordes_y_alineacion(ws, fila_inicio, fila_inicio + len(indicadores) - 1, 1, 2)

    ws.merge_cells('H10:L10')
    ws['H10'] = 'LECTURA EJECUTIVA'
    ws['H10'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    ws['H10'].fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
    ws['H10'].alignment = Alignment(horizontal='center')

    lectura = [
        'El reporte resume los movimientos de stock registrados por el sistema.',
        'Las salidas representan ventas o egresos de mercadería.',
        'Las entradas representan compras, reposiciones o ingresos de proveedor.',
        'La base completa permite trabajar con estadística descriptiva, gráficos, correlación e inferencia.',
        'Los campos stock anterior y stock posterior permiten auditar la variación del inventario.'
    ]

    fila_lectura = 11

    for texto in lectura:
        ws.merge_cells(start_row=fila_lectura, start_column=8, end_row=fila_lectura, end_column=12)
        celda = ws.cell(row=fila_lectura, column=8)
        celda.value = f'• {texto}'
        celda.font = Font(name='Calibri', size=10, color='1E293B')
        celda.fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        celda.alignment = Alignment(vertical='center', wrap_text=True)
        fila_lectura += 1

    ws['A21'] = 'Tipo'
    ws['B21'] = 'Cantidad de movimientos'
    ws['A22'] = 'Salidas'
    ws['B22'] = totales['total_ventas']
    ws['A23'] = 'Entradas'
    ws['B23'] = totales['total_compras']
    ws['A24'] = 'Ajustes'
    ws['B24'] = totales['total_ajustes']

    chart = BarChart()
    chart.title = 'Movimientos por tipo'
    chart.y_axis.title = 'Cantidad'
    chart.x_axis.title = 'Tipo'

    data = Reference(ws, min_col=2, min_row=21, max_row=24)
    cats = Reference(ws, min_col=1, min_row=22, max_row=24)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 13

    ws.add_chart(chart, 'D21')

    top_productos = top_productos_salida(movimientos, 5)

    ws['H21'] = 'Producto'
    ws['I21'] = 'Unidades vendidas'

    for idx, item in enumerate(top_productos, start=22):
        ws.cell(row=idx, column=8).value = item[0]
        ws.cell(row=idx, column=9).value = item[1]

    if top_productos:
        pie = PieChart()
        pie.title = 'Top productos vendidos'

        labels = Reference(ws, min_col=8, min_row=22, max_row=21 + len(top_productos))
        data = Reference(ws, min_col=9, min_row=21, max_row=21 + len(top_productos))

        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.height = 7
        pie.width = 10

        ws.add_chart(pie, 'J21')

    for col in range(1, 13):
        ws.column_dimensions[get_column_letter(col)].width = 18

    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['H'].width = 34
    ws.column_dimensions['I'].width = 20

    configurar_impresion(ws)

    return ws


def crear_hoja_diccionario(wb):
    ws = wb.create_sheet('Diccionario de Datos')
    ws.sheet_properties.tabColor = '64748B'
    ws.sheet_view.showGridLines = False

    ws.merge_cells('A1:D1')
    ws['A1'] = 'DICCIONARIO DE DATOS — EXPORTACIÓN ESTADÍSTICA'
    ws['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='17324D', end_color='17324D', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 32

    encabezados = ['Campo', 'Tipo de dato', 'Descripción', 'Uso estadístico']

    for col, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(row=3, column=col)
        celda.value = encabezado
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        celda.alignment = Alignment(horizontal='center')

    filas = [
        ['id_movimiento', 'Numérico', 'Identificador único del movimiento.', 'Trazabilidad y control de registros.'],
        ['fecha', 'Fecha/Hora', 'Momento en que se registró el movimiento.', 'Análisis por período y tendencia temporal.'],
        ['tipo_movimiento', 'Categórica nominal', 'Entrada, salida o ajuste.', 'Tablas de frecuencia y comparación de grupos.'],
        ['producto', 'Categórica nominal', 'Nombre del producto.', 'Ranking de rotación y análisis por producto.'],
        ['categoria', 'Categórica nominal', 'Rubro o categoría.', 'Distribución por categorías.'],
        ['stock_anterior', 'Cuantitativa discreta', 'Stock antes del movimiento.', 'Control de inventario.'],
        ['cantidad', 'Cuantitativa discreta', 'Cantidad movida.', 'Media, mediana, moda y dispersión.'],
        ['stock_posterior', 'Cuantitativa discreta', 'Stock después del movimiento.', 'Detección de niveles críticos.'],
        ['precio_unitario', 'Cuantitativa continua', 'Precio unitario.', 'Relación precio-demanda.'],
        ['importe_estimado', 'Cuantitativa continua', 'Cantidad por precio.', 'Impacto económico.'],
        ['usuario_responsable', 'Categórica nominal', 'Usuario que registró.', 'Auditoría operativa.'],
        ['motivo', 'Texto', 'Causa del movimiento.', 'Contexto de la operación.']
    ]

    fila_actual = 4

    for fila in filas:
        for col, valor in enumerate(fila, start=1):
            celda = ws.cell(row=fila_actual, column=col)
            celda.value = valor
            celda.alignment = Alignment(vertical='top', wrap_text=True)

        fila_actual += 1

    aplicar_bordes_y_alineacion(ws, 3, fila_actual - 1, 1, 4)
    agregar_tabla_excel(ws, 'Tabla_Diccionario_Datos', 3, 'A', fila_actual - 1, 'D')

    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 46
    ws.column_dimensions['D'].width = 52

    configurar_impresion(ws)

    return ws


def generar_excel_movimientos(movimientos, fecha_desde, fecha_hasta):
    wb = Workbook()

    hoja_inicial = wb.active
    wb.remove(hoja_inicial)

    wb.properties.title = 'Reporte profesional de movimientos de stock'
    wb.properties.subject = 'Exportación estadística generada por Stockeado'
    wb.properties.creator = session.get('usuario_nombre', 'Stockeado')
    wb.properties.description = 'Archivo Excel profesional para análisis estadístico de movimientos de stock.'

    totales = calcular_totales(movimientos)

    crear_hoja_resumen_profesional(wb, movimientos, fecha_desde, fecha_hasta)

    agregar_hoja_detalle(
        wb=wb,
        nombre_hoja='Base Completa',
        movimientos=movimientos,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        color_pestana='2563EB',
        descripcion='Base completa de movimientos exportados para análisis estadístico.'
    )

    agregar_hoja_detalle(
        wb=wb,
        nombre_hoja='Ventas',
        movimientos=totales['ventas'],
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        color_pestana='DC2626',
        descripcion='Movimientos de salida asociados a ventas o egresos de stock.'
    )

    agregar_hoja_detalle(
        wb=wb,
        nombre_hoja='Compras',
        movimientos=totales['compras'],
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        color_pestana='16A34A',
        descripcion='Movimientos de entrada asociados a compras, reposiciones o ingresos de proveedor.'
    )

    crear_hoja_diccionario(wb)

    archivo = BytesIO()
    wb.save(archivo)
    archivo.seek(0)

    nombre_archivo = f'reporte_profesional_stockeado_{fecha_desde}_a_{fecha_hasta}.xlsx'

    return send_file(
        archivo,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=nombre_archivo
    )


# =========================================================
# SEGUNDA MEJORA: NOTIFICACIONES GMAIL POR STOCK BAJO
# =========================================================

def obtener_destinatarios_alerta(conn):
    destinatarios_configurados = os.getenv('STOCK_ALERT_EMAILS', '').strip()

    if destinatarios_configurados:
        return [
            correo.strip()
            for correo in destinatarios_configurados.split(',')
            if correo.strip()
        ]

    cursor = conn.cursor(dictionary=True)

    cursor.execute("""
        SELECT email
        FROM usuarios
        WHERE activo = 1
          AND rol IN ('encargado', 'admin')
          AND email IS NOT NULL
          AND email <> ''
    """)

    usuarios = cursor.fetchall()
    cursor.close()

    return [usuario['email'] for usuario in usuarios]


def registrar_notificacion_stock(conn, id_producto, destinatarios, asunto, cuerpo, modo, estado, detalle_error=''):
    try:
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO notificaciones_stock
            (id_producto, destinatarios, asunto, cuerpo, modo, estado, detalle_error)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """, (
            id_producto,
            ', '.join(destinatarios),
            asunto,
            cuerpo,
            modo,
            estado,
            detalle_error
        ))

        cursor.close()

    except Exception as error:
        print('No se pudo registrar la notificación de stock:', error)


def escribir_log_notificacion(asunto, cuerpo, destinatarios):
    ruta_log = os.path.join(LOG_FOLDER, 'notificaciones_stock.log')

    with open(ruta_log, 'a', encoding='utf-8') as archivo:
        archivo.write('\n' + '=' * 80 + '\n')
        archivo.write(f'FECHA: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}\n')
        archivo.write(f'DESTINATARIOS: {", ".join(destinatarios)}\n')
        archivo.write(f'ASUNTO: {asunto}\n')
        archivo.write('CUERPO:\n')
        archivo.write(cuerpo)
        archivo.write('\n' + '=' * 80 + '\n')


def enviar_gmail(destinatarios, asunto, cuerpo):
    gmail_user = os.getenv('GMAIL_USER', '').strip()
    gmail_app_password = os.getenv('GMAIL_APP_PASSWORD', '').strip().replace(' ', '')
    notificaciones_activas = os.getenv('MAIL_NOTIFICATIONS_ENABLED', 'true').strip().lower()

    if notificaciones_activas != 'true':
        raise ValueError('El envío de correos está desactivado: MAIL_NOTIFICATIONS_ENABLED no está en true.')

    if not gmail_user or not gmail_app_password:
        raise ValueError('Faltan las variables GMAIL_USER o GMAIL_APP_PASSWORD en el archivo .env.')

    mensaje = EmailMessage()
    mensaje['Subject'] = asunto
    mensaje['From'] = gmail_user
    mensaje['To'] = ', '.join(destinatarios)
    mensaje.set_content(cuerpo)

    with smtplib.SMTP('smtp.gmail.com', 587) as servidor:
        servidor.starttls()
        servidor.login(gmail_user, gmail_app_password)
        servidor.send_message(mensaje)


def enviar_notificacion_stock_bajo(conn, producto, nuevo_stock):
    destinatarios = obtener_destinatarios_alerta(conn)

    if not destinatarios:
        destinatarios = ['sin_destinatario_configurado@stockeado.local']

    asunto = f"ALERTA DE STOCK BAJO — {producto['nombre']}"

    cuerpo = f"""
Sistema Stockeado — Alerta automática de inventario

El producto "{producto['nombre']}" alcanzó un nivel crítico de stock.

Datos del producto:
- Producto: {producto['nombre']}
- Stock actual: {nuevo_stock}
- Stock mínimo configurado: {producto['stock_minimo']}
- Precio unitario: ${producto['precio']}
- Fecha de detección: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}

Recomendación:
Se recomienda revisar el inventario y generar una reposición a la brevedad.

Este correo fue generado automáticamente por el sistema Stockeado.
"""

    escribir_log_notificacion(asunto, cuerpo, destinatarios)

    try:
        enviar_gmail(destinatarios, asunto, cuerpo)

        registrar_notificacion_stock(
            conn=conn,
            id_producto=producto['id'],
            destinatarios=destinatarios,
            asunto=asunto,
            cuerpo=cuerpo,
            modo='gmail',
            estado='enviado',
            detalle_error=''
        )

        return True

    except Exception as error:
        registrar_notificacion_stock(
            conn=conn,
            id_producto=producto['id'],
            destinatarios=destinatarios,
            asunto=asunto,
            cuerpo=cuerpo,
            modo='gmail',
            estado='fallido',
            detalle_error=str(error)
        )

        print('No se pudo enviar Gmail de stock bajo:', error)
        return False


def gestionar_alerta_stock(conn, producto, nuevo_stock):
    cursor = conn.cursor(dictionary=True)

    if nuevo_stock <= producto['stock_minimo']:
        cursor.execute("""
            SELECT id
            FROM alertas
            WHERE id_producto = %s
              AND resuelta = 0
            LIMIT 1
        """, (producto['id'],))

        alerta_existente = cursor.fetchone()

        if not alerta_existente:
            mensaje = (
                f"Stock de '{producto['nombre']}' ({nuevo_stock}) "
                f"alcanzó o quedó por debajo del mínimo ({producto['stock_minimo']})."
            )

            cursor_insert = conn.cursor()

            cursor_insert.execute("""
                INSERT INTO alertas (tipo, mensaje, id_producto)
                VALUES (%s, %s, %s)
            """, ('stock_bajo', mensaje, producto['id']))

            cursor_insert.close()

            enviar_notificacion_stock_bajo(conn, producto, nuevo_stock)

    else:
        cursor_update = conn.cursor()

        cursor_update.execute("""
            UPDATE alertas
            SET resuelta = 1
            WHERE id_producto = %s
              AND resuelta = 0
        """, (producto['id'],))

        cursor_update.close()

    cursor.close()


# =========================================================
# LOGIN
# =========================================================

@app.route('/', methods=['GET', 'POST'])
def login():
    error = None

    if request.method == 'POST':
        email = request.form['email']
        password = request.form['password']

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(
            'SELECT * FROM usuarios WHERE email = %s AND password = %s AND activo = 1',
            (email, password)
        )

        usuario = cursor.fetchone()

        cursor.close()
        conn.close()

        if usuario:
            session['usuario_id'] = usuario['id']
            session['usuario_nombre'] = usuario['nombre']
            session['usuario_rol'] = usuario['rol']
            return redirect(url_for('dashboard'))

        error = 'Email o contraseña incorrectos.'

    return render_template('login.html', error=error)


# =========================================================
# DASHBOARD
# =========================================================

@app.route('/dashboard')
def dashboard():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    busqueda = request.args.get('busqueda', '').strip()
    categoria_filtro = request.args.get('categoria', '').strip()

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute('SELECT COUNT(*) AS total FROM productos WHERE activo = 1')
    total_productos = cursor.fetchone()['total']

    cursor.execute('SELECT COUNT(*) AS total FROM alertas WHERE resuelta = 0')
    total_alertas = cursor.fetchone()['total']

    cursor.execute('SELECT COUNT(*) AS total FROM productos WHERE stock_actual <= stock_minimo AND activo = 1')
    total_reponer = cursor.fetchone()['total']

    cursor.execute('SELECT COUNT(*) AS total FROM movimientos_stock WHERE DATE(fecha) = CURDATE()')
    movimientos_hoy = cursor.fetchone()['total']

    cursor.execute('SELECT id, nombre FROM categorias WHERE activo = 1 ORDER BY nombre')
    categorias = cursor.fetchall()

    resultados_busqueda = []
    busqueda_realizada = bool(busqueda or categoria_filtro)

    if busqueda_realizada:
        query = """
            SELECT
                p.*,
                c.nombre AS categoria_nombre
            FROM productos p
            LEFT JOIN categorias c ON p.id_categoria = c.id
            WHERE p.activo = 1
        """

        params = []

        if busqueda:
            query += """
                AND (
                    p.nombre LIKE %s
                    OR p.descripcion LIKE %s
                    OR c.nombre LIKE %s
                )
            """
            params.extend([
                f'%{busqueda}%',
                f'%{busqueda}%',
                f'%{busqueda}%'
            ])

        if categoria_filtro:
            query += " AND p.id_categoria = %s"
            params.append(categoria_filtro)

        query += " ORDER BY p.nombre"

        cursor.execute(query, params)
        resultados_busqueda = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        'dashboard.html',
        nombre=session['usuario_nombre'],
        rol=session['usuario_rol'],
        total_productos=total_productos,
        total_alertas=total_alertas,
        total_reponer=total_reponer,
        movimientos_hoy=movimientos_hoy,
        categorias=categorias,
        busqueda=busqueda,
        categoria_filtro=categoria_filtro,
        resultados_busqueda=resultados_busqueda,
        busqueda_realizada=busqueda_realizada
    )


# =========================================================
# PRODUCTOS
# =========================================================

@app.route('/productos')
def productos():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    busqueda = request.args.get('busqueda', '').strip()
    categoria_filtro = request.args.get('categoria', '').strip()

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    query = """
        SELECT p.*, c.nombre AS categoria_nombre
        FROM productos p
        LEFT JOIN categorias c ON p.id_categoria = c.id
        WHERE p.activo = 1
    """

    params = []

    if busqueda:
        query += " AND (p.nombre LIKE %s OR p.descripcion LIKE %s)"
        params.extend([f'%{busqueda}%', f'%{busqueda}%'])

    if categoria_filtro:
        query += " AND p.id_categoria = %s"
        params.append(categoria_filtro)

    query += " ORDER BY p.nombre"

    cursor.execute(query, params)
    lista_productos = cursor.fetchall()

    cursor.execute('SELECT * FROM categorias WHERE activo = 1 ORDER BY nombre')
    categorias = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        'productos.html',
        productos=lista_productos,
        categorias=categorias,
        nombre=session['usuario_nombre'],
        rol=session['usuario_rol'],
        busqueda=busqueda,
        categoria_filtro=categoria_filtro
    )


@app.route('/productos/registrar', methods=['POST'])
def registrar_producto():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    if session['usuario_rol'] not in ['encargado', 'admin']:
        return redirect(url_for('productos'))

    nombre = request.form['nombre']
    descripcion = request.form.get('descripcion', '')
    precio = request.form['precio']
    stock_actual = int(request.form['stock_actual'])
    stock_minimo = int(request.form['stock_minimo'])
    id_categoria = request.form['id_categoria']

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("""
        INSERT INTO productos (nombre, descripcion, precio, stock_actual, stock_minimo, id_categoria)
        VALUES (%s, %s, %s, %s, %s, %s)
    """, (nombre, descripcion, precio, stock_actual, stock_minimo, id_categoria))

    id_producto = cursor.lastrowid
    cursor.close()

    cursor_producto = conn.cursor(dictionary=True)
    cursor_producto.execute('SELECT * FROM productos WHERE id = %s', (id_producto,))
    producto = cursor_producto.fetchone()
    cursor_producto.close()

    if producto:
        gestionar_alerta_stock(conn, producto, stock_actual)

    conn.commit()
    conn.close()

    return redirect(url_for('productos'))


# =========================================================
# MOVIMIENTOS
# =========================================================

@app.route('/movimientos')
def movimientos():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute('SELECT * FROM productos WHERE activo = 1 ORDER BY nombre')
    lista_productos = cursor.fetchall()

    cursor.execute("""
        SELECT m.*, p.nombre AS producto_nombre, u.nombre AS usuario_nombre
        FROM movimientos_stock m
        JOIN productos p ON m.id_producto = p.id
        JOIN usuarios u ON m.id_usuario = u.id
        ORDER BY m.fecha DESC
        LIMIT 50
    """)

    lista_movimientos = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        'movimientos.html',
        productos=lista_productos,
        movimientos=lista_movimientos,
        nombre=session['usuario_nombre'],
        rol=session['usuario_rol']
    )


@app.route('/movimientos/registrar', methods=['POST'])
def registrar_movimiento():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    id_producto = int(request.form['id_producto'])
    tipo = request.form['tipo']
    cantidad = int(request.form['cantidad'])
    motivo = request.form.get('motivo', '')

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute('SELECT * FROM productos WHERE id = %s', (id_producto,))
    producto = cursor.fetchone()

    if not producto:
        cursor.close()
        conn.close()
        flash('El producto seleccionado no existe.')
        return redirect(url_for('movimientos'))

    if cantidad <= 0:
        cursor.close()
        conn.close()
        flash('La cantidad debe ser mayor a cero.')
        return redirect(url_for('movimientos'))

    stock_anterior = producto['stock_actual']

    if tipo == 'entrada':
        nuevo_stock = stock_anterior + cantidad

    elif tipo == 'salida':
        if cantidad > stock_anterior:
            cursor.close()
            conn.close()
            flash('No se puede registrar una salida mayor al stock disponible.')
            return redirect(url_for('movimientos'))

        nuevo_stock = stock_anterior - cantidad

    elif tipo == 'ajuste':
        nuevo_stock = cantidad

    else:
        cursor.close()
        conn.close()
        flash('Tipo de movimiento inválido.')
        return redirect(url_for('movimientos'))

    cursor2 = conn.cursor()

    cursor2.execute(
        'UPDATE productos SET stock_actual = %s WHERE id = %s',
        (nuevo_stock, id_producto)
    )

    cursor2.execute("""
        INSERT INTO movimientos_stock
        (tipo, cantidad, stock_anterior, stock_posterior, precio_unitario, motivo, id_producto, id_usuario)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
    """, (
        tipo,
        cantidad,
        stock_anterior,
        nuevo_stock,
        producto['precio'],
        motivo,
        id_producto,
        session['usuario_id']
    ))

    cursor2.close()

    producto['stock_actual'] = nuevo_stock
    gestionar_alerta_stock(conn, producto, nuevo_stock)

    conn.commit()

    cursor.close()
    conn.close()

    flash('Movimiento registrado correctamente.')
    return redirect(url_for('movimientos'))


# =========================================================
# IMPORTAR EXCEL DE PROVEEDOR
# =========================================================

@app.route('/importar', methods=['GET', 'POST'])
def importar():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    if session['usuario_rol'] not in ['encargado', 'admin']:
        return redirect(url_for('dashboard'))

    resultados = []
    errores = []
    productos_con_similares = []
    archivo_procesado = False

    if request.method == 'POST' and 'archivo' in request.files:
        archivo = request.files.get('archivo')

        if not archivo or not archivo.filename.endswith('.xlsx'):
            errores.append('El archivo debe ser un Excel (.xlsx).')
        else:
            ruta = os.path.join(UPLOAD_FOLDER, archivo.filename)
            archivo.save(ruta)

            try:
                wb = openpyxl.load_workbook(ruta)
                ws = wb.active
                filas = list(ws.iter_rows(min_row=2, values_only=True))

                conn = get_db_connection()
                cursor = conn.cursor(dictionary=True)
                cursor2 = conn.cursor()

                for fila in filas:
                    if not any(fila):
                        continue

                    proveedor = str(fila[0]) if len(fila) > 0 and fila[0] else 'Desconocido'
                    fecha_excel = str(fila[1]) if len(fila) > 1 and fila[1] else ''
                    nombre_producto = str(fila[2]) if len(fila) > 2 and fila[2] else ''
                    cantidad = int(fila[3]) if len(fila) > 3 and fila[3] else 0

                    if not nombre_producto or cantidad <= 0:
                        continue

                    cursor.execute(
                        'SELECT * FROM productos WHERE LOWER(nombre) = LOWER(%s) AND activo = 1',
                        (nombre_producto,)
                    )

                    producto = cursor.fetchone()

                    if producto:
                        stock_anterior = producto['stock_actual']
                        nuevo_stock = stock_anterior + cantidad

                        cursor2.execute(
                            'UPDATE productos SET stock_actual = %s WHERE id = %s',
                            (nuevo_stock, producto['id'])
                        )

                        motivo = f"Ingreso por proveedor: {proveedor} — fecha factura: {fecha_excel}"

                        cursor2.execute("""
                            INSERT INTO movimientos_stock
                            (tipo, cantidad, stock_anterior, stock_posterior, precio_unitario, motivo, id_producto, id_usuario)
                            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                        """, (
                            'entrada',
                            cantidad,
                            stock_anterior,
                            nuevo_stock,
                            producto['precio'],
                            motivo,
                            producto['id'],
                            session['usuario_id']
                        ))

                        producto['stock_actual'] = nuevo_stock
                        gestionar_alerta_stock(conn, producto, nuevo_stock)

                        resultados.append({
                            'producto': producto['nombre'],
                            'proveedor': proveedor,
                            'cantidad': cantidad,
                            'stock_nuevo': nuevo_stock,
                            'ok': True
                        })

                    else:
                        productos_con_similares.append({
                            'nombre': nombre_producto,
                            'proveedor': proveedor,
                            'fecha': fecha_excel,
                            'cantidad': cantidad,
                            'similares': []
                        })

                conn.commit()
                cursor.close()
                cursor2.close()
                conn.close()

                archivo_procesado = True

            except Exception as error:
                errores.append(f'Error al procesar el archivo: {str(error)}')

            finally:
                if os.path.exists(ruta):
                    os.remove(ruta)

    return render_template(
        'importar.html',
        nombre=session['usuario_nombre'],
        rol=session['usuario_rol'],
        resultados=resultados,
        errores=errores,
        productos_con_similares=productos_con_similares,
        archivo_procesado=archivo_procesado
    )


# =========================================================
# EXPORTACIÓN ESTADÍSTICA
# =========================================================

@app.route('/exportar-movimientos')
def exportar_movimientos():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    if not usuario_puede_exportar():
        return redirect(url_for('dashboard'))

    hoy = datetime.now().date()
    primer_dia_mes = hoy.replace(day=1)

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute("""
        SELECT id, nombre
        FROM categorias
        WHERE activo = 1
        ORDER BY nombre
    """)

    categorias = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        'exportar_movimientos.html',
        nombre=session['usuario_nombre'],
        rol=session['usuario_rol'],
        categorias=categorias,
        fecha_desde=primer_dia_mes.strftime('%Y-%m-%d'),
        fecha_hasta=hoy.strftime('%Y-%m-%d')
    )


@app.route('/exportar-movimientos/descargar', methods=['POST'])
def descargar_movimientos_estadistica():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    if not usuario_puede_exportar():
        return redirect(url_for('dashboard'))

    fecha_desde = request.form.get('fecha_desde')
    fecha_hasta = request.form.get('fecha_hasta')
    tipo_movimiento = request.form.get('tipo_movimiento', 'todos')
    id_categoria = request.form.get('id_categoria', '')
    formato = request.form.get('formato', 'xlsx')

    fechas_validas, mensaje_error = validar_rango_fechas(fecha_desde, fecha_hasta)

    if not fechas_validas:
        flash(mensaje_error)
        return redirect(url_for('exportar_movimientos'))

    movimientos = consultar_movimientos_exportacion(
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        tipo_movimiento=tipo_movimiento,
        id_categoria=id_categoria
    )

    if not movimientos:
        flash('No existen movimientos registrados para el rango y filtros seleccionados.')
        return redirect(url_for('exportar_movimientos'))

    registrar_auditoria_exportacion(
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        tipo_movimiento=tipo_movimiento,
        id_categoria=id_categoria,
        formato=formato,
        cantidad_registros=len(movimientos)
    )

    if formato == 'csv':
        return generar_csv_movimientos(movimientos, fecha_desde, fecha_hasta)

    return generar_excel_movimientos(movimientos, fecha_desde, fecha_hasta)


# =========================================================
# ALERTAS
# =========================================================

@app.route('/alertas')
def alertas():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute("""
        SELECT a.*, p.nombre AS producto_nombre, p.stock_actual, p.stock_minimo
        FROM alertas a
        JOIN productos p ON a.id_producto = p.id
        WHERE a.resuelta = 0
        ORDER BY a.fecha DESC
    """)

    lista_alertas = cursor.fetchall()

    cursor.execute("""
        SELECT p.*, c.nombre AS categoria_nombre
        FROM productos p
        LEFT JOIN categorias c ON p.id_categoria = c.id
        WHERE p.stock_actual <= p.stock_minimo
          AND p.activo = 1
        ORDER BY (p.stock_minimo - p.stock_actual) DESC
    """)

    productos_reponer = cursor.fetchall()

    cursor.execute("""
        SELECT n.*, p.nombre AS producto_nombre
        FROM notificaciones_stock n
        JOIN productos p ON n.id_producto = p.id
        ORDER BY n.fecha_envio DESC
        LIMIT 20
    """)

    notificaciones = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        'alertas.html',
        alertas=lista_alertas,
        productos_reponer=productos_reponer,
        notificaciones=notificaciones,
        nombre=session['usuario_nombre'],
        rol=session['usuario_rol']
    )


@app.route('/alertas/resolver/<int:id>')
def resolver_alerta(id):
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    if session['usuario_rol'] not in ['encargado', 'admin']:
        return redirect(url_for('alertas'))

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute('UPDATE alertas SET resuelta = 1 WHERE id = %s', (id,))

    conn.commit()
    cursor.close()
    conn.close()

    return redirect(url_for('alertas'))


# =========================================================
# USUARIOS
# =========================================================

@app.route('/usuarios')
def gestion_usuarios():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    if session.get('usuario_rol') != 'admin':
        return redirect(url_for('dashboard'))

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute('SELECT * FROM usuarios ORDER BY nombre')
    usuarios = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        'usuarios.html',
        usuarios=usuarios,
        nombre=session['usuario_nombre'],
        rol=session['usuario_rol']
    )


@app.route('/usuarios/registrar', methods=['POST'])
def registrar_usuario():
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    if session.get('usuario_rol') != 'admin':
        return redirect(url_for('dashboard'))

    nombre = request.form['nombre']
    email = request.form['email']
    password = request.form['password']
    rol = request.form['rol']

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("""
        INSERT INTO usuarios (nombre, email, password, rol, activo)
        VALUES (%s, %s, %s, %s, 1)
    """, (nombre, email, password, rol))

    conn.commit()
    cursor.close()
    conn.close()

    flash('Usuario registrado correctamente.')
    return redirect(url_for('gestion_usuarios'))


@app.route('/usuarios/toggle/<int:id>')
def toggle_usuario(id):
    if 'usuario_id' not in session:
        return redirect(url_for('login'))

    if session.get('usuario_rol') != 'admin':
        return redirect(url_for('dashboard'))

    if id == session.get('usuario_id'):
        flash('No podés desactivar tu propio usuario.')
        return redirect(url_for('gestion_usuarios'))

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("""
        UPDATE usuarios
        SET activo = CASE WHEN activo = 1 THEN 0 ELSE 1 END
        WHERE id = %s
    """, (id,))

    conn.commit()
    cursor.close()
    conn.close()

    return redirect(url_for('gestion_usuarios'))


# =========================================================
# LOGOUT
# =========================================================

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


if __name__ == '__main__':
    app.run(debug=True)